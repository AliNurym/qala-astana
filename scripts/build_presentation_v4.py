import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
PPTX_PATH = OUTPUT_DIR / "Qala_Astana_v4.pptx"

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
except ImportError:
    print("python-pptx not available yet")
    sys.exit(1)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_BG = RGBColor(11, 19, 36)
PAPER = RGBColor(255, 255, 255)
CYAN = RGBColor(0, 242, 254)
BLUE = RGBColor(79, 172, 254)
GREEN = RGBColor(56, 239, 125)
GOLD = RGBColor(255, 209, 102)
MUTED = RGBColor(155, 176, 211)
CARD_BG = RGBColor(21, 35, 65)

def apply_slide_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG

def add_header(slide, title_text, category_text, slide_num):
    # Category badge
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = CYAN

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(10), Inches(0.8))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PAPER

    # Slide Number
    num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.5), Inches(1), Inches(0.4))
    tf_n = num_box.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = f"{slide_num:02d} / 07"
    p_n.font.size = Pt(12)
    p_n.font.color.rgb = MUTED
    p_n.alignment = PP_ALIGN.RIGHT

# -------------------------------------------------------------
# SLIDE 1: Cover
# -------------------------------------------------------------
s1 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s1)

# Hackathon Badge
b_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(10), Inches(0.5))
p = b_box.text_frame.paragraphs[0]
p.text = "BUILD WITH AI · ТРЕК 1: НИЦ SMART CITY & GDG ASTANA"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = CYAN

# Title
main_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.2))
tf = main_box.text_frame
p = tf.paragraphs[0]
p.text = "QALA ASTANA 2.0"
p.font.size = Pt(56)
p.font.bold = True
p.font.color.rgb = PAPER

p2 = tf.add_paragraph()
p2.text = "Интеллектуальный ИИ-помощник по прогулкам в городе Астана"
p2.font.size = Pt(24)
p2.font.color.rgb = GREEN

# Description box
desc_box = s1.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.5), Inches(2.2))
tf_d = desc_box.text_frame
tf_d.word_wrap = True
p = tf_d.paragraphs[0]
p.text = "Рабочая RAG-система реального времени, которая подсказывает пользователю, куда сходить и что сделать поблизости с учетом локации, времени, компании, доступной среды и городской афиши."
p.font.size = Pt(16)
p.font.color.rgb = MUTED

p_sub = tf_d.add_paragraph()
p_sub.text = "✨ Реализованы все 5 стретч-целей: Мини-маршруты · Мультиязычность KZ/RU/EN · Доступная среда · Групповое планирование · Афиша событий"
p_sub.font.size = Pt(14)
p_sub.font.bold = True
p_sub.font.color.rgb = GOLD

# -------------------------------------------------------------
# SLIDE 2: Problem & 5 Hackathon Scenarios
# -------------------------------------------------------------
s2 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s2)
add_header(s2, "Пользовательские сценарии и задача из ТЗ", "Концепция и сценарии", 2)

body_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Проблема: когда есть 30–60 минут, искать места по разным картам неудобно, а стандартные LLM галлюцинируют."
p.font.size = Pt(16)
p.font.color.rgb = MUTED

cases = [
    ("1. Прогулка у Mega Silk Way (30–60 мин)", "Подбор ТРЦ Mega Silk Way, Линейного парка Экспо и сферы Нұр Әлем с оценкой пути пешком."),
    ("2. Тихое кафе с розеткой и Wi-Fi", "Верификация розеток, Wi-Fi и спокойной атмосферы для работы (Кафе «Көктем»)."),
    ("3. С ребёнком (площадки / интерактивы)", "Интерактивный парк Happylon Ocean Park, детские городки и 8 этажей экспозиций Нұр Әлем."),
    ("4. Турист с 2 часами перед вылетом (Аэропорт)", "Честное информирование о границах покрытия Астаны без ложных галлюцинаций."),
    ("5. Необычная ночная активность", "Видовая терраса на крыше отеля Hilton (The Cloud Bar) до 02:00, ночные пабы и вечерняя подсветка."),
]

for title, desc in cases:
    p_t = tf.add_paragraph()
    p_t.text = f"• {title}: "
    p_t.font.size = Pt(15)
    p_t.font.bold = True
    p_t.font.color.rgb = CYAN
    
    # Append desc
    run = p_t.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = PAPER

# -------------------------------------------------------------
# SLIDE 3: Two-stage RAG Pipeline
# -------------------------------------------------------------
s3 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s3)
add_header(s3, "Архитектура двухэтапного RAG-подхода", "Техническая архитектура", 3)

body_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

blocks = [
    ("ЭТАП 1: RETRIEVAL (Умный сбор и семантический поиск)", [
        "Автоматический сбор Overpass API (OSM) + Обогащение из реестра Smart City: 821 проверенный POI.",
        "Мультиязычные поля: названия и описания на казахском (KK), русском (RU), английском (EN).",
        "Гибридный BM25 + Intent matching: учет интентов («кофе», «детям», «ночь», «группа», «коляска»).",
        "Пространственный расчет Haversine + матрица пеших маршрутов OpenRouteService (ORS).",
    ]),
    ("ЭТАП 2: GENERATION (Генерация обоснований и плана действий)", [
        "Поддержка Gemini 2.5 Flash, DeepSeek Flash и встроенного Intelligent Local RAG Engine.",
        "Strict Grounding: Модель опирается исключительно на поля POI и не галлюцинирует.",
        "Генерация двух обязательных блоков: «Почему подходит» (мотивация) и «Что сделать там» (план).",
    ]),
]

for title, items in blocks:
    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.size = Pt(16)
    p_t.font.bold = True
    p_t.font.color.rgb = GREEN
    
    for item in items:
        p_i = tf.add_paragraph()
        p_i.text = f"  - {item}"
        p_i.font.size = Pt(14)
        p_i.font.color.rgb = PAPER

# -------------------------------------------------------------
# SLIDE 4: Response Contract & Honest AI
# -------------------------------------------------------------
s4 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s4)
add_header(s4, "Контракт ответа и достоверность информации", "Формат вывода и честность", 4)

body_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

rows = [
    ("1–3 конкретные рекомендации", "Оптимальный объем без информационной перегрузки; ранжирование по релевантности."),
    ("Рейтинг и отзывы", "Звездочный рейтинг (⭐ 4.8) и объем подтвержденных отзывов для доверия."),
    ("Время в пути и время на месте", "Раздельный показ: пешком (ORS или оценка) + ориентир длительности визита по категории."),
    ("Почему подходит + Что сделать там", "Краткая мотивация с учетом контекста + конкретный сценарий действий на месте."),
    ("Актуальная городская афиша", "Интеграция временных выставок («Зеленый прорыв 2030»), световых шоу и живой музыки."),
    ("Честная маркировка и часы", "Часы работы сопровождаются напоминанием «уточните перед визитом», исключая ложные обещания."),
]

for title, desc in rows:
    p = tf.add_paragraph()
    p.text = f"✔ {title}: "
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN
    r = p.add_run()
    r.text = desc
    r.font.bold = False
    r.font.color.rgb = PAPER

# -------------------------------------------------------------
# SLIDE 5: 5 Stretch Goals
# -------------------------------------------------------------
s5 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s5)
add_header(s5, "Реализация всех 5 стретч-целей хакатона", "Баллы за креативность", 5)

body_box = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

goals = [
    ("🚶 Мини-маршрут (цепочка 3 мест)", "Построение логической последовательности (Старт -> Кульминация -> Финал). Расчет сегментов пути, времени на местах, общего хронометража и отрисовка линии на карте!"),
    ("🌐 Мультиязычность (KZ / RU / EN)", "Полная локализация интерфейса и генерация ответов на казахском, русском и английском языках."),
    ("♿ Доступная среда (Wheelchair)", "Фильтрация для маломобильных граждан: отбор объектов с подтвержденными пандусами и лифтами."),
    ("👥 Групповое планирование", "Специальный режим для компаний из 3–5 человек: просторные посадки, командные активности."),
    ("🎪 Городские события и афиша", "Интеграция временных выставок и вечерних шоу (Нұр Әлем, Байтерек, Hilton Rooftop)."),
]

for title, desc in goals:
    p = tf.add_paragraph()
    p.text = f"★ {title}: "
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GOLD
    r = p.add_run()
    r.text = desc
    r.font.bold = False
    r.font.color.rgb = PAPER

# -------------------------------------------------------------
# SLIDE 6: Privacy & Ethics
# -------------------------------------------------------------
s6 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s6)
add_header(s6, "Этика, Privacy и защита данных", "Безопасность пользователей", 6)

body_box = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

items = [
    ("Ноль хранения координат", "Геолокация пользователя обрабатывается исключительно в оперативной памяти во время запроса. Координаты не записываются в базу данных, файлы или серверные логи."),
    ("Изоляция внешних LLM", "Во внешние модели (Gemini / DeepSeek) передаются только текст запроса и обезличенные описания POI. Координаты пользователя в сторонние сервисы ИИ не отправляются."),
    ("Прозрачная атрибуция", "Пользователь всегда видит источник данных: что является верифицированным фактом из реестра POI, а что — сгенерированной идеей модели."),
    ("Защита от Prompt Injection", "Все пользовательские данные и тексты POI изолированы как недоверенный ввод в строгом системном промпте."),
]

for title, desc in items:
    p = tf.add_paragraph()
    p.text = f"🛡️ {title}: "
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN
    r = p.add_run()
    r.text = desc
    r.font.bold = False
    r.font.color.rgb = PAPER

# -------------------------------------------------------------
# SLIDE 7: Conclusion & Roadmap
# -------------------------------------------------------------
s7 = prs.slides.add_slide(prs.slide_layouts[6])
apply_slide_bg(s7)
add_header(s7, "Итоги, готовность и дорожная карта", "План развития", 7)

body_box = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = body_box.text_frame
tf.word_wrap = True

p_curr = tf.paragraphs[0]
p_curr.text = "ТЕКУЩИЙ СТАТУС ПРОЕКТА:"
p_curr.font.size = Pt(16)
p_curr.font.bold = True
p_curr.font.color.rgb = CYAN

curr_items = [
    "База данных: 821 POI Астаны (центр, набережная, Expo, Mega Silk Way).",
    "Прототип: интерактивное веб-приложение с картой, быстрыми кейсами и режимом мини-маршрутов.",
    "Качество и тесты: 23 автоматических unit-теста пройдены на 100%.",
]
for ci in curr_items:
    p = tf.add_paragraph()
    p.text = f"  ✔ {ci}"
    p.font.size = Pt(14)
    p.font.color.rgb = PAPER

p_next = tf.add_paragraph()
p_next.text = "ДОРОЖНАЯ КАРТА РАЗВИТИЯ:"
p_next.font.size = Pt(16)
p_next.font.bold = True
p_next.font.color.rgb = GOLD

next_items = [
    "1. Расширение географии: включение правого берега Астаны и зоны аэропорта.",
    "2. Интеграция с CTS: учет времени автобусов и будущей линии Astana LRT в пеших маршрутах.",
    "3. Партнерство с Astana Tourism и акиматом для автоматической синхронизации событий и фестивалей.",
]
for ni in next_items:
    p = tf.add_paragraph()
    p.text = f"  🚀 {ni}"
    p.font.size = Pt(14)
    p.font.color.rgb = PAPER

prs.save(str(PPTX_PATH))
print(f"Презентация успешно сохранена: {PPTX_PATH}")
